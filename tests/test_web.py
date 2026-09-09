import base64
import http.client
import io
import json
import tempfile
import unittest
import zipfile
from http import HTTPStatus
from pathlib import Path
from unittest.mock import patch


class WebTests(unittest.TestCase):
    def test_chapter_draft_includes_previous_episode_context(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        prompts: list[str] = []

        class FakeClient:
            def complete_json(self, system_prompt: str, user_prompt: str) -> dict[str, str]:
                prompts.append(user_prompt)
                return {
                    "title": "第二集·回信",
                    "outline": "主角追查上一集留下的信件。",
                    "content": "第2集\n△ 主角重新展开上一集的信。",
                }

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("chapter-continuation", "", "secret123")
            project = store.create_project(
                user["id"],
                {"title": "雨夜来信", "genre": "悬疑", "style": "二维国漫"},
            )
            first = store.create_chapter(user["id"], project["id"], {"title": "第一集"})
            store.update_chapter(
                user["id"],
                project["id"],
                first["id"],
                {"outline": "主角收到未来来信。", "content": "第1集\n△ 主角拆开未来来信。"},
            )
            second = store.create_chapter(user["id"], project["id"], {"title": "第二集"})
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.offline_demo = False
            handler._build_client = lambda payload: FakeClient()
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_chapter_draft(
                user["id"],
                project["id"],
                second["id"],
                {
                    "brief": "继续调查寄信人",
                    "progress_id": "chapter-draft-progress",
                },
            )

            self.assertIn("上一集连续性上下文", prompts[0])
            self.assertIn("主角收到未来来信", prompts[0])
            self.assertIn("主角拆开未来来信", prompts[0])
            self.assertEqual(response["chapter"]["episode_no"], 2)
            progress = handler._get_quick_progress("drama", "chapter-draft-progress")
            self.assertEqual(progress["status"], "completed")
            self.assertEqual(progress["percent"], 100)

    def test_quick_create_generates_script_and_seedance_storyboard_in_one_call(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("quick-create", "", "secret123")
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.offline_demo = True
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_quick_create(
                user["id"],
                {"title": "雨夜来信", "brief": "失忆快递员收到一封来自未来的信"},
            )

            project = response["project"]
            chapter = response["chapter"]
            production = response["production"]
            self.assertEqual(len(project["chapters"]), 1)
            self.assertIn("△ ", chapter["content"])
            self.assertTrue(production["characters"])
            self.assertTrue(production["scenes"])
            self.assertTrue(production["shots"])
            self.assertIn("Seedance 2.0", production["shots"][0]["video_prompt"])
            self.assertEqual(len(store.list_projects(user["id"])), 1)

    def test_novel_generate_distills_full_writing_package(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("novel-writer", "", "secret123")
            project = store.create_project(
                user["id"],
                {
                    "title": "饺子网文",
                    "product_type": "novel",
                    "description": "重生逆袭",
                },
            )
            chapter = store.create_chapter(
                user["id"], project["id"], {"title": "第 1 章", "content": "主角重生。"}
            )
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.offline_demo = True
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_novel_generate(
                user["id"],
                project["id"],
                chapter["id"],
                {
                    "source_text": "落魄程序员重生高考前。",
                    "progress_id": "novel-package-progress",
                },
            )

            package = response["novel_package"]
            saved = response["chapter"]
            self.assertEqual(package["type"], "novel_package")
            self.assertIn("读者契约", package["stages"])
            self.assertIn("前三章雷达", package["stages"])
            self.assertEqual(package["metadata"]["source"], "jiaozi-novel-engine-v1")
            self.assertTrue(package["characters"])
            self.assertIn("reader_contract", package)
            self.assertIn("quality_gate", package)
            self.assertIn("chapter_text", package)
            self.assertEqual(saved["production"]["type"], "novel_package")
            self.assertEqual(store.list_projects(user["id"])[0]["novel_stage_count"], 6)
            progress = handler._get_quick_progress("novel", "novel-package-progress")
            self.assertEqual(progress["status"], "completed")
            self.assertEqual(progress["percent"], 100)

    def test_jubensha_generate_distills_round_table_package(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("jubensha-writer", "", "secret123")
            project = store.create_project(
                user["id"],
                {
                    "title": "钟楼旧宴",
                    "product_type": "jubensha",
                    "description": "六个旧友在闭馆钟楼重聚。",
                },
            )
            chapter = store.create_chapter(
                user["id"], project["id"], {"title": "第 1 幕", "content": "圆桌上出现死者邀请函。"}
            )
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.offline_demo = True
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_jubensha_generate(
                user["id"],
                project["id"],
                chapter["id"],
                {
                    "source_text": "六个旧友在闭馆钟楼重聚，桌上出现一封写给死者的邀请函。",
                    "player_count": "5人",
                    "duration": "3小时",
                    "progress_id": "jubensha-package-progress",
                },
            )

            package = response["jubensha_package"]
            saved = response["chapter"]
            self.assertEqual(package["type"], "jubensha_package")
            self.assertIn("席位契约", package["stages"])
            self.assertIn("真相骨架", package["stages"])
            self.assertEqual(package["metadata"]["source"], "jiaozi-jubensha-engine-v1")
            self.assertTrue(package["roles"])
            self.assertTrue(package["clues"])
            self.assertTrue(package["rounds"])
            self.assertIn("dm_manual", package)
            self.assertIn("playability_gate", package)
            self.assertEqual(package["play_contract"]["player_count"], "5人")
            self.assertEqual(package["play_contract"]["duration"], "3小时")
            self.assertEqual(saved["production"]["type"], "jubensha_package")
            summary = store.list_projects(user["id"])[0]
            self.assertGreaterEqual(summary["jubensha_role_count"], 1)
            self.assertGreaterEqual(summary["jubensha_clue_count"], 1)
            self.assertGreaterEqual(summary["jubensha_round_count"], 1)
            progress = handler._get_quick_progress("jubensha", "jubensha-package-progress")
            self.assertEqual(progress["status"], "completed")
            self.assertEqual(progress["percent"], 100)

    def test_drama_generate_records_package_progress(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        class FakeProduction:
            metadata = {"generation_mode": "test"}

            @staticmethod
            def to_dict() -> dict[str, object]:
                return {"title": "测试制作包", "characters": [], "scenes": [], "shots": []}

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("drama-package", "", "secret123")
            project = store.create_project(user["id"], {"title": "测试短剧"})
            chapter = store.create_chapter(user["id"], project["id"], {"title": "第一集"})
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler._run_generation = lambda payload: (FakeProduction(), {})
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_project_generate(
                user["id"],
                project["id"],
                chapter["id"],
                {"progress_id": "drama-package-progress"},
            )

            self.assertEqual(response["chapter"]["production"]["title"], "测试制作包")
            progress = handler._get_quick_progress("drama", "drama-package-progress")
            self.assertEqual(progress["status"], "completed")
            self.assertEqual(progress["percent"], 100)

    def test_quick_progress_is_isolated_by_user_and_records_completion(self) -> None:
        from ai_drama_agent.web import DramaWebHandler

        DramaWebHandler._set_quick_progress(
            "novel",
            "progress-test",
            "user-1",
            100,
            "创作完成",
            "已保存",
            status="completed",
        )

        progress = DramaWebHandler._get_quick_progress("novel", "progress-test")
        self.assertEqual(progress["user_id"], "user-1")
        self.assertEqual(progress["scope"], "novel")
        self.assertEqual(progress["status"], "completed")
        self.assertEqual(progress["percent"], 100)

    def test_novel_quick_create_uses_novel_namespace(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("novel-quick-create", "", "secret123")
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.offline_demo = True
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_novel_quick_create(
                user["id"],
                {
                    "title": "万里归途",
                    "brief": "都市重生开局。",
                    "genre": "都市",
                    "style": "重生",
                    "progress_id": "novel-progress",
                },
            )

            project = response["project"]
            chapter = response["chapter"]
            package = response["novel_package"]
            self.assertEqual(project["product_type"], "novel")
            self.assertEqual(package["type"], "novel_package")
            self.assertTrue(chapter["production"])
            progress = handler._get_quick_progress("novel", "novel-progress")
            self.assertEqual(progress["status"], "completed")

    def test_jubensha_quick_create_uses_jubensha_namespace(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("jubensha-quick-create", "", "secret123")
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.offline_demo = True
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler._handle_jubensha_quick_create(
                user["id"],
                {
                    "title": "钟楼旧宴",
                    "brief": "六个旧友在闭馆钟楼重聚。",
                    "genre": "还原本",
                    "style": "现代",
                    "player_count": "6人",
                    "duration": "4小时",
                    "progress_id": "jubensha-progress",
                },
            )

            project = response["project"]
            chapter = response["chapter"]
            package = response["jubensha_package"]
            self.assertEqual(project["product_type"], "jubensha")
            self.assertEqual(package["type"], "jubensha_package")
            self.assertTrue(chapter["production"])
            progress = handler._get_quick_progress("jubensha", "jubensha-progress")
            self.assertEqual(progress["status"], "completed")

    def test_production_package_patch_updates_novel_cards(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("novel-editor", "", "secret123")
            project = store.create_project(
                user["id"],
                {"title": "饺子网文", "product_type": "novel", "description": "重生逆袭"},
            )
            chapter = store.create_chapter(
                user["id"], project["id"], {"title": "第 1 章", "content": "主角重生。"}
            )
            store.save_production(
                user["id"],
                project["id"],
                chapter["id"],
                {
                    "type": "novel_package",
                    "title": "旧书名",
                    "positioning": "旧卖点",
                    "reader_contract": {"target_reader": "老读者"},
                    "world": {"background": "旧背景"},
                    "outline": {"chapter_beats": [{"title": "旧节点"}]},
                    "characters": [{"name": "旧角色"}],
                    "chapter_title": "第 1 章",
                    "chapter_outline": "旧章纲",
                    "chapter_text": "旧正文",
                },
            )
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler._require_user = lambda: user
            response: dict[str, object] = {}
            handler._send_json = response.update
            handler.path = f"/api/projects/{project['id']}/chapters/{chapter['id']}/production"
            handler._read_json = lambda: {
                "production": {
                    "type": "novel_package",
                    "title": "新书名",
                    "positioning": "新卖点",
                    "reader_contract": {"target_reader": "新读者"},
                    "world": {"background": "新背景"},
                    "outline": {"chapter_beats": [{"title": "新节点"}]},
                    "characters": [{"name": "新角色"}],
                    "chapter_title": "第 1 章",
                    "chapter_outline": "新章纲",
                    "chapter_text": "新正文",
                }
            }

            handler.do_PATCH()

            self.assertEqual(response["chapter"]["production"]["title"], "新书名")
            self.assertEqual(response["chapter"]["production"]["characters"][0]["name"], "新角色")


    def test_delete_chapter_endpoint_removes_saved_assets(self) -> None:
        from ai_drama_agent.store import LocalStore
        from ai_drama_agent.web import DramaWebHandler

        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(Path(directory) / "app_state.json")
            user = store.register("delete-chapter", "", "secret123")
            project = store.create_project(user["id"], {"title": "删除章节"})
            chapter = store.create_chapter(user["id"], project["id"], {"title": "第一集"})
            store.save_production(
                user["id"],
                project["id"],
                chapter["id"],
                {
                    "characters": [{"id": "CHAR-001"}],
                    "scenes": [{"id": "SCENE-001"}],
                    "props": [{"id": "PROP-001"}],
                    "shots": [{"id": "SHOT-001"}],
                },
            )
            handler = object.__new__(DramaWebHandler)
            handler.store = store
            handler.path = f"/api/projects/{project['id']}/chapters/{chapter['id']}"
            handler._require_user = lambda: user
            response: dict[str, object] = {}
            handler._send_json = response.update

            handler.do_DELETE()

            self.assertEqual(response, {"ok": True})
            self.assertEqual(store.get_project(user["id"], project["id"])["chapters"], [])

    def test_api_helpers(self) -> None:
        from ai_drama_agent.web import _optional_text, _positive_int, _required_text

        self.assertEqual(_required_text({"title": "测试"}, "title"), "测试")
        with self.assertRaises(ValueError):
            _required_text({"title": "   "}, "title")
        self.assertEqual(_optional_text({}, "style", "默认风格"), "默认风格")
        self.assertEqual(_positive_int("24", 12), 24)
        self.assertEqual(_positive_int("-1", 12), 12)

    def test_base64_upload_body_limit_accepts_maximum_document(self) -> None:
        from ai_drama_agent.web import MAX_BODY_BYTES, MAX_UPLOAD_BYTES

        encoded_size = ((MAX_UPLOAD_BYTES + 2) // 3) * 4
        self.assertGreater(MAX_BODY_BYTES, encoded_size)

    def test_error_payload_shape(self) -> None:
        payload = json.loads(json.dumps({"error": "字段 title 不能为空。"}, ensure_ascii=False))
        self.assertEqual(payload["error"], "字段 title 不能为空。")
        self.assertEqual(HTTPStatus.BAD_REQUEST, 400)

    def test_store_error_status_distinguishes_validation_from_missing_data(self) -> None:
        from ai_drama_agent.store import StoreError
        from ai_drama_agent.web import _store_error_status

        self.assertEqual(_store_error_status(StoreError("用户名至少需要 2 个字符。")), HTTPStatus.BAD_REQUEST)
        self.assertEqual(_store_error_status(StoreError("项目不存在或无权访问。")), HTTPStatus.NOT_FOUND)

    def test_rule_based_mode_label_value(self) -> None:
        payload = json.loads(
            json.dumps({"mode": "offline-rule-based"}, ensure_ascii=False)
        )
        self.assertEqual(payload["mode"], "offline-rule-based")

    def test_real_model_requires_explicit_api_key(self) -> None:
        from ai_drama_agent.llm import OpenAICompatibleClient

        with self.assertRaisesRegex(RuntimeError, "尚未配置真实模型 API Key"):
            OpenAICompatibleClient.from_config("", "https://api.openai.com/v1", "gpt-4.1")

    def test_reasoning_effort_is_added_only_when_explicit(self) -> None:
        from ai_drama_agent.llm import OpenAICompatibleClient

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *args: object) -> None:
                return None

            def read(self) -> bytes:
                return b'{"choices":[{"message":{"content":"{}"}}]}'

        captured: list[dict[str, object]] = []
        captured_request_headers: list[dict[str, str]] = []

        def fake_urlopen(request, timeout):
            captured.append(json.loads(request.data.decode("utf-8")))
            captured_request_headers.append(dict(request.header_items()))
            return FakeResponse()

        with patch("urllib.request.urlopen", side_effect=fake_urlopen):
            OpenAICompatibleClient.from_config("key", "https://example.com/v1", "o3", "high").complete_json("s", "u")
            OpenAICompatibleClient.from_config("key", "https://example.com/v1", "gpt-4.1", "auto").complete_json("s", "u")

        self.assertEqual(captured[0]["reasoning_effort"], "high")
        self.assertNotIn("reasoning_effort", captured[1])
        normalized_headers = {key.lower(): value for key, value in captured_request_headers[0].items()}
        self.assertIn("user-agent", normalized_headers)
        self.assertIn("accept", normalized_headers)

    def test_model_http_error_is_user_readable(self) -> None:
        from ai_drama_agent.llm import _format_http_error

        message = _format_http_error(401, '{"error":{"message":"Incorrect API key"}}')
        self.assertIn("Incorrect API key", message)
        self.assertIn("请检查 API Key", message)

    def test_cloudflare_1010_is_reported_as_upstream_access_denied(self) -> None:
        from ai_drama_agent.llm import _format_http_error

        detail = """
        <!doctype html>
        <html>
        <head><title>Access denied | api.qlhazycoder.top used Cloudflare to restrict access</title></head>
        <body>
        <p>The owner of this website (api.qlhazycoder.top) has banned your access based on your browser's signature.</p>
        <h1>Error 1010</h1>
        </body>
        </html>
        """.strip()
        message = _format_http_error(403, detail)
        self.assertIn("Cloudflare", message)
        self.assertIn("api.qlhazycoder.top", message)
        self.assertIn("Error 1010", message)
        self.assertIn("上游风控或封禁", message)
        self.assertNotIn("当前 API Key 没有访问该模型或接口的权限", message)
        self.assertNotIn("<html>", message)

    def test_proxy_message_with_html_cause_is_compacted(self) -> None:
        from ai_drama_agent.llm import _format_http_error

        detail = (
            "CC Switch local proxy failed while handling Codex endpoint /responses. "
            "Provider: My Codex; model: codex-auto-review; upstream_status: HTTP 403; "
            "cause: <!doctype html><html><head><title>Access denied | api.qlhazycoder.top used "
            "Cloudflare to restrict access</title></head><body><p>The owner of this website "
            "(api.qlhazycoder.top) has banned your access based on your browser's signature.</p>"
            "<h1>Error 1010</h1></body></html>"
        )
        message = _format_http_error(403, detail)
        self.assertIn("CC Switch local proxy failed", message)
        self.assertIn("Cloudflare", message)
        self.assertIn("api.qlhazycoder.top", message)
        self.assertLess(len(message), 420)

    def test_model_http_error_hides_oversized_provider_payload(self) -> None:
        from ai_drama_agent.llm import _format_http_error

        message = _format_http_error(500, "x" * 800)
        self.assertLess(len(message), 560)
        self.assertTrue(message.endswith("..."))

    def test_remote_disconnect_is_user_readable(self) -> None:
        from ai_drama_agent.llm import OpenAICompatibleClient

        with patch("urllib.request.urlopen", side_effect=http.client.RemoteDisconnected("closed")):
            with self.assertRaisesRegex(RuntimeError, "远端提前断开"):
                OpenAICompatibleClient.from_config(
                    "key",
                    "https://example.com/v1",
                    "gpt-4.1",
                ).complete_json("s", "u")

    def test_import_txt_script(self) -> None:
        from ai_drama_agent.web import _extract_uploaded_script

        payload = {
            "filename": "scene.txt",
            "content_base64": base64.b64encode("雨夜，林默推门而入。".encode("utf-8")).decode("ascii"),
        }
        self.assertEqual(_extract_uploaded_script(payload), "雨夜，林默推门而入。")

    def test_import_docx_script(self) -> None:
        from ai_drama_agent.web import _extract_uploaded_script

        xml = """
        <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body>
            <w:p><w:r><w:t>第一幕</w:t></w:r></w:p>
            <w:p><w:r><w:t>角色走进房间。</w:t></w:r></w:p>
          </w:body>
        </w:document>
        """.strip()
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", xml)

        payload = {
            "filename": "scene.docx",
            "content_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
        }
        self.assertEqual(_extract_uploaded_script(payload), "第一幕\n角色走进房间。")

    def test_prompt_document_contains_all_shot_prompt_types(self) -> None:
        from docx import Document

        from ai_drama_agent.web import _build_prompt_document

        content = _build_prompt_document(
            {"title": "雨夜归来", "style": "电影感国漫"},
            {"episode_no": 1, "title": "便利店", "production": {}},
            {
                "characters": [{"name": "林默", "turnaround_prompt": "角色提示词"}],
                "scenes": [{"name": "便利店", "environment_prompt": "场景提示词"}],
                "props": [{"name": "雨伞", "description": "道具提示词"}],
                "shots": [
                    {
                        "id": "shot_001",
                        "shot_size": "近景",
                        "camera_position": "平视",
                        "duration_seconds": 3,
                        "first_frame_prompt": "首帧提示词",
                        "video_prompt": "视频提示词",
                        "last_frame_prompt": "尾帧提示词",
                        "negative_prompt": "负面提示词",
                    }
                ],
            },
        )

        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for expected in ("角色提示词", "场景提示词", "道具提示词", "首帧提示词", "视频提示词", "尾帧提示词", "负面提示词"):
            self.assertIn(expected, text)

    def test_prompt_document_exports_canonical_material_map_prompts(self) -> None:
        from docx import Document

        from ai_drama_agent.web import _build_prompt_document

        content = _build_prompt_document(
            {"title": "资产一致性", "style": "电影感国漫"},
            {"episode_no": 1, "title": "第一集"},
            {
                "characters": [{"name": "林默", "turnaround_prompt": "旧的角色简写"}],
                "scenes": [{"name": "办公室", "environment_prompt": "旧的场景简写"}],
                "props": [{"name": "信封", "description": "旧的道具简写"}],
                "material_map": [
                    {"id": "C001", "type": "角色资产", "prompt": "资产编号：C001。完整角色提示词"},
                    {"id": "S001", "type": "场景资产", "prompt": "资产编号：S001。完整场景提示词"},
                    {"id": "P001", "type": "道具资产", "prompt": "资产编号：P001。完整道具提示词"},
                ],
                "shots": [],
            },
        )
        text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
        for expected in ("C001", "S001", "P001", "完整角色提示词", "完整场景提示词", "完整道具提示词"):
            self.assertIn(expected, text)
        for stale in ("旧的角色简写", "旧的场景简写", "旧的道具简写"):
            self.assertNotIn(stale, text)

    def test_chat_completion_text_supports_common_relay_shapes(self) -> None:
        from ai_drama_agent.llm import _extract_chat_completion_text

        self.assertEqual(
            _extract_chat_completion_text({"choices": [{"message": {"content": "plain"}}]}),
            "plain",
        )
        self.assertEqual(
            _extract_chat_completion_text({"choices": [{"message": {"content": [{"text": "part one"}, {"text": "part two"}]}}]}),
            "part one\npart two",
        )
        self.assertEqual(
            _extract_chat_completion_text({"choices": [{"text": "legacy"}]}),
            "legacy",
        )
        self.assertEqual(_extract_chat_completion_text({"output_text": "fallback"}), "fallback")
        with self.assertRaisesRegex(RuntimeError, "choices"):
            _extract_chat_completion_text({})

    def test_script_document_contains_outline_and_content(self) -> None:
        from docx import Document

        from ai_drama_agent.web import _build_script_document

        content = _build_script_document(
            {"title": "Project Title"},
            {"title": "Episode One", "outline": "Outline text", "content": "Scene one\nDialogue two"},
        )
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for expected in ("Project Title", "Episode One", "Outline text", "Scene one", "Dialogue two"):
            self.assertIn(expected, text)

    def test_production_archive_contains_script_and_storyboard_documents(self) -> None:
        from docx import Document

        from ai_drama_agent.web import _build_production_archive

        content = _build_production_archive(
            {"title": "雨夜归来", "style": "电影感国漫"},
            {"episode_no": 1, "title": "便利店", "outline": "本章梗概", "content": "第一场\n角色对白"},
            {
                "characters": [],
                "scenes": [],
                "props": [],
                "shots": [{"id": "shot_001", "video_prompt": "分镜视频提示词"}],
            },
        )
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), 2)
            self.assertTrue(any(name.endswith("-剧本.docx") for name in names))
            self.assertTrue(any(name.endswith("-分镜提示词.docx") for name in names))
            script_name = next(name for name in names if name.endswith("-剧本.docx"))
            prompt_name = next(name for name in names if name.endswith("-分镜提示词.docx"))
            script_text = "\n".join(
                paragraph.text
                for paragraph in Document(io.BytesIO(archive.read(script_name))).paragraphs
            )
            prompt_text = "\n".join(
                paragraph.text
                for paragraph in Document(io.BytesIO(archive.read(prompt_name))).paragraphs
            )
        self.assertIn("角色对白", script_text)
        self.assertIn("分镜视频提示词", prompt_text)

    def test_novel_archive_contains_summary_and_chapter_document(self) -> None:
        from docx import Document

        from ai_drama_agent.web import _build_production_archive

        content = _build_production_archive(
            {"title": "重生之夜", "product_type": "novel"},
            {
                "title": "第 1 章 开局",
                "outline": "主角重生回到高考前。",
                "content": "第一章正文\n主角重新开始。",
            },
            {
                "type": "novel_package",
                "title": "重生之夜",
                "positioning": "都市重生逆袭",
                "stages": ["读者契约", "前三章雷达"],
                "reader_contract": {"target_reader": "移动端追读用户"},
                "opening_radar": {"hook": "重生回到高考前"},
                "world": {"rule": "成绩可以改变命运"},
                "outline": {"volume": "主线"},
                "serial_plan": {"cadence": "日更"},
                "publish_gate": {"word_count": 2800},
                "chapter_title": "第 1 章 开局",
                "chapter_text": "第一章正文\n主角重新开始。",
            },
        )
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            summary_name = next(name for name in names if name.endswith("-连载方案.md"))
            chapter_name = next(name for name in names if name.endswith("-首章正文.docx"))
            self.assertTrue(any(name.endswith("-项目.json") for name in names))
            summary_text = archive.read(summary_name).decode("utf-8")
            chapter_text = "\n".join(
                paragraph.text
                for paragraph in Document(io.BytesIO(archive.read(chapter_name))).paragraphs
            )
        self.assertIn("读者契约", summary_text)
        self.assertIn("首章正文预览", summary_text)
        self.assertIn("主角重新开始", chapter_text)

    def test_jubensha_archive_contains_summary_and_opening_document(self) -> None:
        from docx import Document

        from ai_drama_agent.web import _build_production_archive

        content = _build_production_archive(
            {"title": "钟楼旧宴", "product_type": "jubensha"},
            {
                "title": "第 1 幕 入席",
                "outline": "所有人回到钟楼。",
                "content": "DM开场词\n今晚的圆桌已准备好。",
            },
            {
                "type": "jubensha_package",
                "title": "钟楼旧宴",
                "positioning": "六人还原本",
                "stages": ["席位契约", "真相骨架"],
                "play_contract": {"player_count": "6人"},
                "truth_spine": {"opening_question": "谁最早进入钟楼"},
                "roles": [{"name": "沈砚"}],
                "clues": [{"id": "CLUE-001", "truth_target": "时间异常"}],
                "rounds": [{"name": "入席开局"}],
                "dm_manual": {"opening": "DM开场词"},
                "playability_gate": {"role_balance": "平衡"},
                "chapter_title": "第 1 幕 入席",
                "opening_script": "DM开场词\n今晚的圆桌已准备好。",
            },
        )
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            summary_name = next(name for name in names if name.endswith("-剧本杀制作包.md"))
            chapter_name = next(name for name in names if name.endswith("-开局正文.docx"))
            self.assertTrue(any(name.endswith("-项目.json") for name in names))
            summary_text = archive.read(summary_name).decode("utf-8")
            chapter_text = "\n".join(
                paragraph.text
                for paragraph in Document(io.BytesIO(archive.read(chapter_name))).paragraphs
            )
        self.assertIn("席位契约", summary_text)
        self.assertIn("真相骨架", summary_text)
        self.assertIn("今晚的圆桌已准备好", chapter_text)

    def test_pdf_always_mode_skips_native_text_layer(self) -> None:
        from ai_drama_agent.web import _extract_uploaded_script

        payload = {
            "filename": "scene.pdf",
            "ocr_mode": "always",
            "content_base64": base64.b64encode(b"fake-pdf").decode("ascii"),
        }
        with patch("ai_drama_agent.web._extract_pdf_text", return_value="OCR 剧本") as extractor:
            self.assertEqual(_extract_uploaded_script(payload), "OCR 剧本")
        extractor.assert_called_once_with(b"fake-pdf", mode="always")


if __name__ == "__main__":
    unittest.main()
